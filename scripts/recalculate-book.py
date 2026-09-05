"""Read-only independent arithmetic audit; never writes the workbook."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('workbook', type=Path)
parser.add_argument('output', type=Path)
args = parser.parse_args()
wb = load_workbook(args.workbook, data_only=True, read_only=True)
d = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb.worksheets}
wb.close()
checks = []
result = {}


def check(name, computed, expected, tolerance=1e-7):
    checks.append({'name': name, 'computed': computed, 'cached': expected, 'difference': computed - expected, 'pass': abs(computed - expected) < tolerance})


for name, start, assumptions in [('悲观', 9, 41), ('基准', 31, 42), ('乐观', 53, 43)]:
    wacc = d['假设'][f'B{assumptions}']
    growth = d['假设'][f'C{assumptions}']
    roic = d['假设'][f'D{assumptions}']
    nopat = 604.8
    series = []
    for index, row in enumerate(range(start, start + 12)):
        nopat *= 1 + d['DCF'][f'B{row}']
        series.append({'year': 2031 + index, 'nopat': nopat, 'roic': d['DCF'][f'E{row}']})
    for index, item in enumerate(series):
        row = start + index
        nopat = item['nopat']
        investment = ((series[index + 1]['nopat'] - nopat) / series[index + 1]['roic']
                      if index < 11 else nopat * growth / roic)
        fcff = nopat - investment
        pv = fcff / (1 + wacc) ** (index + 1)
        for col, value in [('C', nopat), ('F', investment), ('G', fcff), ('I', pv)]:
            check(f'DCF {col}{row}', value, d['DCF'][f'{col}{row}'])
        if index < 11:
            check(f'{name} forward timing {row}', investment * series[index + 1]['roic'], series[index + 1]['nopat'] - nopat)
        item.update(reinvestment=investment, fcff=fcff, pv=pv)
    explicit = sum(item['pv'] for item in series)
    terminal_nopat = nopat * (1 + growth)
    terminal_investment = terminal_nopat * growth / roic
    terminal_fcff = terminal_nopat - terminal_investment
    terminal_pv = terminal_fcff / (wacc - growth) / (1 + wacc) ** 12
    ev = explicit + terminal_pv
    price = (ev - 500 + 1030.1) / 300
    check(f'{name} price', price, d['DCF'][f'L{start + 12}'])
    check(f'{name} stable investment rate', terminal_investment / terminal_nopat, growth / roic)
    check(f'{name} 2042 transition', series[-1]['reinvestment'] * roic, terminal_nopat - nopat)
    check(f'{name} 2043 forward timing', terminal_investment * roic, terminal_nopat * growth)
    result[name] = {'series': series, 'explicitPV': explicit, 'terminalNOPAT': terminal_nopat,
                    'terminalReinvestment': terminal_investment, 'terminalFCFF': terminal_fcff,
                    'terminalPV': terminal_pv, 'EV': ev, 'price': price, 'terminalShare': terminal_pv / ev}

base = result['基准']
sensitivity = []
for row in range(8, 14):
    growth = d['敏感性'][f'A{row}']
    values = []
    for col in 'BCDEFGH':
        wacc = d['敏感性'][f'{col}7']
        nopat = base['series'][-1]['nopat']
        last_fcff = nopat * (1 - growth / .09)
        terminal_fcff = nopat * (1 + growth) * (1 - growth / .09)
        explicit = sum(item['fcff'] / (1 + wacc) ** (item['year'] - 2030) for item in base['series'][:-1])
        price = (explicit + last_fcff / (1 + wacc) ** 12 + terminal_fcff / (wacc - growth) / (1 + wacc) ** 12 + 530.1) / 300
        check(f'敏感性 {col}{row}', price, d['敏感性'][f'{col}{row}'])
        if abs(wacc - .09) < 1e-9:
            check(f'growth neutral g={growth}', price, base['price'])
        values.append(price)
    sensitivity.append({'g': growth, 'prices': values})
result['sensitivity'] = sensitivity

single = d['单店估值']
for row in [9, 10, 11]:
    price = single[f'B{row}'] * (1 + single[f'C{row}']) / (single[f'D{row}'] - single[f'C{row}'])
    check(f'single store {row}', price, single[f'F{row}'])

history = d['历史经营']
for col in 'BCDEFGH':
    assets = sum(history[f'{col}{row}'] for row in [21, 22, 23, 24, 25, 26])
    liabilities_equity = sum(history[f'{col}{row}'] for row in [28, 29, 30, 31])
    check(f'{col} balance sheet', assets, liabilities_equity)
    nwc = sum(history[f'{col}{row}'] for row in [22, 23, 24]) - sum(history[f'{col}{row}'] for row in [28, 29])
    check(f'{col} NWC', nwc, history[f'{col}33'])
    cfo = history[f'{col}17'] + history[f'{col}12'] - history[f'{col}34']
    check(f'{col} CFO or labeled 2024 proxy', cfo, history[f'{col}42'])
    fcff = history[f'{col}15'] + history[f'{col}12'] - history[f'{col}43'] - history[f'{col}34']
    check(f'{col} FCFF', fcff, history[f'{col}47'])
    check(f'{col} FCFE', cfo - history[f'{col}43'] + history[f'{col}48'], history[f'{col}50'])
    if col != 'B':
        prev = chr(ord(col) - 1)
        cash = history[f'{prev}21'] + cfo - history[f'{col}43'] - history[f'{col}45'] + history[f'{col}48'] + history[f'{col}49']
        check(f'{col} cash roll', cash, history[f'{col}21'], .05)
        check(f'{col} PPE roll', history[f'{prev}25'] + history[f'{col}43'] - history[f'{col}12'], history[f'{col}25'])

cohort = d['门店批次']
for row in range(40, 46):
    for col in 'IJKLMN':
        year = cohort[f'{col}39']
        start = cohort[f'A{row}']
        factor = 0 if year < start else [.6, .8, 1.0][min(year - start, 2)]
        check(f'cohort {col}{row}', cohort[f'B{row}'] * cohort[f'C{row}'] * factor, cohort[f'{col}{row}'])
    roic = cohort[f'D{row}'] * (1 - cohort['B35']) / (cohort['B32'] + cohort['B33'])
    check(f'cohort ROIC {row}', roic, cohort[f'G{row}'])
for col, table_col in zip('BCDEFG', 'IJKLMN'):
    check(f'cohort total {col}', 120 + sum(cohort[f'{table_col}{row}'] for row in range(40, 46)), cohort[f'{col}12'])
    check(f'cohort revenue bridge {col}', cohort[f'{col}12'] + cohort[f'{col}13'], cohort[f'{col}9'])

result['singleStore'] = {'normalProfit': 15 - 5 - 1.5 + .8 - .2, 'proxyCFO': 9.1 + 3 - 4,
                         'proxyFCF': 9.1 + 3 - 4 - 11, 'ownersCash': 9.1 + 3 - 1 - 4,
                         'ROE': 9.1 / ((40 + 49.1) / 2), 'ROIC': 9.6 / 48}
result['pricing'] = {}
for quantity in [.97, .92]:
    revenue = 5200 * 1.05 * quantity + 400
    cost = 1600 * 1.08 * quantity + 276
    result['pricing'][str(quantity)] = {'revenue': revenue, 'totalGrowth': revenue / 5600 - 1,
                                      'directGrowth': 1.05 * quantity - 1, 'cost': cost, 'EBIT': revenue - cost - 2668 - 300}
result['stress'] = {}
for key, revenue_factor, cost_factor, expense_factor in [('mild', .9, .92, .98), ('severe', .8, .85, .96)]:
    ebit = 5600 * revenue_factor - 1876 * cost_factor - 2668 * expense_factor - 300
    pretax = ebit - 50.4
    net = pretax - max(pretax, 0) * .2
    result['stress'][key] = {'EBIT': ebit, 'NI': net, 'CFO': net + 300, 'interestCover': ebit / 50.4}
operating = result['stress']['severe']['CFO'] - 4.8
opening = 1198.1 - 168
result['stress']['cashCalendar'] = {'openingAvailable': opening, 'operatingAfterAllInterest': operating,
                                   'annualFlowPaused': operating - 100 - 80,
                                   'annualFlowFullCapex': operating - 850 - 80,
                                   'knownClosingPaused': opening + operating - 100 - 80,
                                   'knownClosingFullCapex': opening + operating - 850 - 80,
                                   'existingDebtMaturities': 'unknown', 'monthlyTiming': 'unknown'}
check('per-share weighted EPS', 571.2 / 262.5, 2.176)
check('per-share ending denominator', 571.2 / 300, 1.904)
check('owner cash excludes growth NWC', result['singleStore']['ownersCash'], 7.1)
check('cost transmission revenue basis', 5200 * .05, 260)
check('cost transmission cost basis', 1600 * .08, 128)
check('coupon price unchanged', 21 - 3, 20 - 2)
scale = (45 * 300 + 500 - 1030.1) / base['EV']
check('reverse DCF', (base['EV'] * scale - 500 + 1030.1) / 300, 45)
result['reverseDCF'] = {'scale': scale, 'firstNOPAT': base['series'][0]['nopat'] * scale,
                        'firstGrowth': base['series'][0]['nopat'] * scale / 604.8 - 1}
result['checks'] = checks
result['failures'] = [check for check in checks if not check['pass']]
args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'comparisons': len(checks), 'failures': len(result['failures'])}))
if result['failures']:
    raise SystemExit(1)
